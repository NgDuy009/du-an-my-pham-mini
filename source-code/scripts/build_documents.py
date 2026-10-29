"""Generate editable report and presentation templates for the project."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'documents'
DOCS.mkdir(exist_ok=True)
SECTIONS = [
('Tóm tắt', 'Lumière Cosmetic là ứng dụng web quản lý cửa hàng mỹ phẩm mini. Hệ thống minh họa kiến trúc web nhiều tầng gồm giao diện người dùng, backend REST API, cơ sở dữ liệu quan hệ PostgreSQL và MinIO Object Storage. Các thành phần được đóng gói và kết nối bằng Docker Compose; database và object storage sử dụng named volume để dữ liệu không mất khi container bị tạo lại. Đồ án tập trung vào nghiệp vụ quản lý sản phẩm, lô hàng, khách hàng, đơn hàng, tệp đính kèm và báo cáo cơ bản.'),
('1. Giới thiệu bài toán', 'Cửa hàng mỹ phẩm nhỏ thường quản lý danh mục, nhãn hàng, nhà cung cấp, khách hàng, tồn kho và đơn hàng bằng sổ ghi chép hoặc nhiều bảng tính tách rời. Cách làm này gây khó khăn khi tra cứu mã sản phẩm, kiểm tra lô hàng và hạn sử dụng, tổng hợp doanh thu, cũng như lưu trữ ảnh sản phẩm và chứng từ. Đồ án xây dựng một hệ thống quản lý thống nhất, có thể khởi chạy trên máy cá nhân để trình diễn kiến thức điện toán đám mây.'),
('2. Mục tiêu hệ thống', 'Mục tiêu chính là xây dựng giao diện có thể sử dụng trên trình duyệt; triển khai REST API; lưu dữ liệu có quan hệ bằng SQL; lưu file qua MinIO; quản lý quyền Administrator/Staff; và chứng minh tính bền vững của dữ liệu bằng Docker volumes. Mục tiêu mở rộng gồm quản lý lô/hạn sử dụng, giao dịch bán hàng có transaction và báo cáo doanh thu.'),
('3. Phạm vi', 'Phạm vi gồm danh mục, sản phẩm, nhãn hàng, nhà cung cấp, khách hàng, lô hàng, đơn hàng, chi tiết đơn, tệp upload, tài khoản, audit log và báo cáo cơ bản. Không nằm trong phạm vi bắt buộc: thanh toán online, nhiều chi nhánh, tích hợp ERP, Kubernetes, cloud công cộng, high availability hoặc microservices. Đây là ứng dụng học tập, không thay thế phần mềm kế toán hoặc hệ thống bán lẻ thương mại đã được kiểm định.'),
('4. Yêu cầu chức năng', 'Người dùng đăng nhập và đăng xuất; xem dashboard; tìm kiếm dữ liệu theo mã/tên; tạo, xem, cập nhật và xóa các thực thể được phân quyền; nhập lô hàng và theo dõi hạn dùng; tạo đơn hàng và in hóa đơn; upload, xem và tải tệp; xem báo cáo; Administrator quản lý tài khoản và nhật ký. Các biểu mẫu có nhãn, validation cơ bản và thông báo kết quả. Các thao tác quan trọng được kiểm tra lại ở backend, không chỉ dựa vào giao diện.'),
('5. Yêu cầu phi chức năng', 'Ứng dụng cần dễ chạy trên Windows với Docker Desktop, có cấu hình thông qua biến môi trường, hạn chế lộ thông tin nhạy cảm, dùng parameterized SQL, hash mật khẩu, xử lý lỗi có thông điệp phù hợp và giao diện co giãn theo kích thước màn hình. Persistence được kiểm chứng bằng việc dừng và chạy lại Compose mà không xóa volume. Cấu hình demo phải dễ tái tạo và tài liệu phải chỉ rõ lệnh khởi động, URL, tài khoản và cách kiểm thử.'),
('6. Use Case', 'Hai actor chính là Administrator và Staff. Cả hai có thể đăng nhập, xem dashboard, tra cứu dữ liệu, thực hiện nghiệp vụ được cấp quyền, bán hàng và quản lý file theo quyền. Administrator có thêm quyền quản lý tài khoản, dữ liệu danh mục và nhật ký. Sơ đồ nguồn được đặt tại docs/diagrams/use-case.mmd. Khi hoàn thiện báo cáo, nhóm nên render sơ đồ Mermaid thành ảnh PNG để đưa vào tài liệu và slide.'),
('7. Kiến trúc tổng thể', 'Lựa chọn phương án giao diện tích hợp backend: Node.js/Express phục vụ EJS và REST API trong cùng container. PostgreSQL lưu dữ liệu quan hệ; MinIO lưu nội dung file; Docker Compose khai báo service, environment, network và volumes. Browser chỉ giao tiếp HTTP với ứng dụng và MinIO Console dành cho demo, không kết nối trực tiếp PostgreSQL. Kiến trúc này phù hợp phương án một Dockerfile ứng dụng được nêu trong tài liệu yêu cầu.'),
('8. Thiết kế database', 'Cơ sở dữ liệu có các bảng users, categories, brands, suppliers, products, product_batches, customers, promotions, orders, order_items, stock_movements, file_metadata và audit_logs. Khóa chính dùng BIGSERIAL; khóa ngoại thể hiện quan hệ; unique constraint chống trùng mã; check constraint ngăn giá/số lượng âm; index hỗ trợ tìm kiếm và tra cứu lịch sử. Chi tiết được định nghĩa trong database/schema.sql; dữ liệu mẫu nằm trong database/seed.sql.'),
('9. Quản lý sản phẩm', 'Mỗi sản phẩm có SKU, tên, danh mục, nhãn hàng, nhà cung cấp, mô tả, thành phần, quy cách, giá nhập, giá bán và ngưỡng cảnh báo tồn. Danh sách cho phép tìm theo SKU hoặc tên; trang chi tiết hiển thị thông tin bản ghi. Giá được lưu dưới dạng NUMERIC để tránh sai số nhị phân của số thực. Danh mục, nhãn hàng và nhà cung cấp được tách thành thực thể riêng để giảm trùng lặp dữ liệu.'),
('10. Quản lý lô và hạn sử dụng', 'Một sản phẩm có thể có nhiều lô nhập. Mỗi lô giữ mã lô, ngày sản xuất, ngày hết hạn, ngày nhận, số lượng nhập, số lượng còn lại và giá nhập. Việc quản lý theo lô giúp xác định sản phẩm sắp hết hạn. Khi bán, hệ thống ưu tiên lô có hạn dùng gần hơn và chỉ sử dụng lô còn hàng, chưa hết hạn. Lịch sử stock_movements ghi nhận nhập, bán và điều chỉnh.'),
('11. Nghiệp vụ đơn hàng', 'Đơn hàng lưu mã đơn, khách hàng, nhân viên tạo, trạng thái, tạm tính, giảm giá, tổng tiền, phương thức thanh toán và thời điểm hoàn tất. Mỗi chi tiết lưu snapshot SKU/tên, số lượng, giá bán và giá vốn tại thời điểm bán. Việc khóa dòng và cập nhật lô trong transaction giảm nguy cơ bán vượt tồn khi hai thao tác diễn ra đồng thời. Nếu tồn không đủ, lỗi được ném để rollback các thay đổi của đơn.'),
('12. Upload và Object Storage', 'Tệp được gửi qua multipart/form-data. Backend giới hạn dung lượng, kiểm tra phần mở rộng và MIME type, sinh object key UUID tránh đè tên, rồi lưu nội dung vào bucket MinIO. PostgreSQL chỉ giữ metadata gồm tên gốc, key, MIME, size, bucket, loại/ID thực thể, người upload và thời gian. Nếu ghi metadata thất bại sau khi upload, ứng dụng cố gắng xóa object vừa tạo để tránh file mồ côi.'),
('13. REST API và HTTP status', 'API cung cấp CRUD cho sản phẩm, danh mục, nhãn hàng, nhà cung cấp và khách hàng; dashboard API trả thống kê; file upload nhận multipart. Các status thường dùng gồm 200, 201, 204, 400, 401, 403, 404 và 500. Tài liệu OpenAPI và Postman Collection được đặt trong thư mục postman. Các request cần session đăng nhập; thao tác quản trị cần role Administrator.'),
('14. Xác thực và phân quyền', 'Ứng dụng dùng session cookie và bcrypt hash. Tài khoản mẫu được tạo lần đầu dựa trên biến môi trường. Administrator được phép quản lý tài khoản và thao tác quản trị; Staff được dùng các màn hình nghiệp vụ đã cho phép. Mật khẩu không được lưu plaintext. Trong ứng dụng thực tế, cần bổ sung CSRF protection, rate limiting, session store bền vững, chính sách mật khẩu và quy trình đổi mật khẩu.'),
('15. Thiết kế giao diện', 'Giao diện dùng sidebar, topbar, dashboard cards, bảng dữ liệu, biểu mẫu, trang chi tiết và thông báo. CSS tự thiết kế với màu trung tính, sắc hồng đất và xanh dịu phù hợp chủ đề mỹ phẩm. Trên màn hình nhỏ, sidebar có thể mở/đóng và bảng có vùng cuộn ngang. Các trạng thái rỗng, lỗi và thành công được thể hiện riêng để người dùng dễ hiểu.'),
('16. Dockerfile và Compose', 'Dockerfile dùng Node.js Alpine, cài dependency production, copy source và khởi chạy ứng dụng. Compose có ba service: application, database và minio. cosmetic-network là bridge network riêng. database-data và minio-data là named volumes. Các thông số nhạy cảm được đọc từ .env, trong khi .env.example chỉ cung cấp tên biến và giá trị mẫu. PostgreSQL được healthcheck trước khi ứng dụng khởi động.'),
('17. Liên hệ VPC/subnet', 'Docker bridge network tạo vùng giao tiếp nội bộ cho các container. Ứng dụng gọi database bằng DNS service database và gọi MinIO bằng minio, không cần IP tĩnh. Mục tiêu cô lập có điểm tương đồng với VPC/subnet nhưng không thể xem Docker network là VPC hoàn chỉnh: nó không thay thế route table, security group, IAM hay các dịch vụ mạng cloud. Việc phân biệt này cần được nêu rõ trong phần lý thuyết.'),
('18. Kiểm thử', 'Cần kiểm thử đăng nhập, phân quyền, CRUD, tìm kiếm, tạo đơn, từ chối bán vượt tồn, upload file hợp lệ/không hợp lệ, xem/tải file, xuất CSV và persistence. Chạy unit test bằng npm test, kiểm tra cú pháp bằng npm run lint:syntax, kiểm tra container bằng docker compose ps và endpoint /health. Ảnh chụp và kết quả thật phải được nhóm bổ sung sau khi chạy trên Docker Desktop; không ghi kết quả giả định thành kết quả đã kiểm chứng.'),
('19. Kiểm chứng persistence', 'Sau khi tạo một bản ghi và upload file, chạy docker compose down, sau đó docker compose up -d. Bản ghi phải còn trong giao diện; file phải còn trong danh sách và có thể tải lại; object vẫn xuất hiện trong MinIO Console. Không dùng docker compose down -v trong bước này vì lệnh đó xóa volumes. Nên ghi lại ảnh chụp trước và sau khi khởi động lại để làm minh chứng báo cáo.'),
('20. Hướng dẫn demo', 'Trình bày theo thứ tự: bài toán; kiến trúc; khởi chạy Compose; đăng nhập; CRUD tối thiểu ba thực thể; tìm kiếm; tạo lô; bán đơn và in hóa đơn; upload/download; Swagger/Postman; mở MinIO Console; trình bày ERD và Deployment Diagram; cuối cùng chứng minh persistence. Tất cả thành viên nên hiểu phần mình phụ trách và lý do sử dụng volume/network. Thời lượng theo PDF là 10–12 phút demo và 5–7 phút hỏi đáp.'),
('21. Hạn chế', 'Phiên bản đồ án chưa xử lý đầy đủ đổi trả, hủy đơn, khuyến mãi nâng cao, nhiều chi nhánh, kế toán, backup tự động hoặc triển khai công khai. MinIO đang dùng root credential trong mạng demo; production cần user/policy tối thiểu quyền. Session store mặc định không phù hợp nhiều replica. Đây là các giới hạn cần nêu trung thực khi bảo vệ, đồng thời tránh trình bày dự án như một hệ thống thương mại hoàn chỉnh.'),
('22. Hướng phát triển và kết luận', 'Có thể mở rộng kiểm thử tích hợp, CSRF, rate limiting, audit đầy đủ hơn, nhập/xuất CSV sản phẩm, hoàn đơn, quản lý chương trình khuyến mãi, báo cáo lợi nhuận, backup định kỳ và triển khai lên VM cloud nếu được yêu cầu. Đồ án minh họa được kiến trúc nhiều tầng, SQL quan hệ, Object Storage, Dockerfile, Compose, biến môi trường và persistence. Trước khi nộp cần xác nhận chạy thành công trên máy nhóm và bổ sung minh chứng thật.'),
('Phụ lục A — Checklist nghiệm thu', 'Kiểm tra docker compose up -d --build; đăng nhập đúng/sai; quyền Administrator và Staff; CRUD tối thiểu ba thực thể; tìm kiếm; upload/download; tệp hiện trong MinIO; Swagger/Postman; tạo đơn và trừ kho; xử lý lỗi khi tồn không đủ; docker compose down rồi up -d vẫn giữ dữ liệu; source không chứa .env hoặc secret; báo cáo và slide có tên thành viên, phân công và ảnh chụp thực tế.'),
('Tài liệu tham khảo', '1. Tài liệu yêu cầu đồ án cuối kỳ học phần Điện toán đám mây do giảng viên cung cấp. 2. Docker Docs — Docker Compose: https://docs.docker.com/compose/. 3. PostgreSQL Documentation: https://www.postgresql.org/docs/. 4. MinIO Documentation: https://min.io/docs/. 5. Express Documentation: https://expressjs.com/.')
]

# Editable DOCX
report = Document()
sec = report.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
styles = report.styles
styles['Normal'].font.name = 'Arial'; styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size, color in [('Title', 28, 'A65F61'), ('Heading 1', 19, 'A65F61'), ('Heading 2', 14, '6B5650')]:
    styles[name].font.name = 'Arial'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].font.bold = True
p = report.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('LUMIÈRE COSMETIC'); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(30); r.font.color.rgb = RGBColor(166,95,97)
for text, size in [('BÁO CÁO ĐỒ ÁN CUỐI KỲ', 18), ('HỌC PHẦN: ĐIỆN TOÁN ĐÁM MÂY', 12)]:
    p = report.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = p.add_run(text); rr.bold = True; rr.font.size = Pt(size)
report.add_paragraph('\n')
for line in ['Đề tài: Website quản lý mỹ phẩm mini đóng gói bằng Docker Compose', 'Sinh viên: [Bổ sung họ tên thành viên]', 'Mã sinh viên: [Bổ sung]', 'Lớp: [Bổ sung]', 'Giảng viên hướng dẫn: [Bổ sung]', 'Đơn vị: [Bổ sung trường / khoa]', 'Năm học: 2026']:
    p = report.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(line)
report.add_page_break()
report.add_heading('Lưu ý hoàn thiện trước khi nộp', level=1)
report.add_paragraph('Đây là báo cáo mẫu được tạo cùng source code. Nhóm cần bổ sung thông tin thành viên, phân công công việc, ảnh chụp giao diện thực tế, kết quả kiểm thử trên máy có Docker Desktop và chỉnh sửa nội dung theo phần đã triển khai. Không xem các mô tả thiết kế là minh chứng rằng hệ thống đã chạy trên máy của nhóm.')
report.add_paragraph('Các sơ đồ Mermaid có source tại docs/diagrams. Có thể render thành PNG/SVG bằng công cụ Mermaid phù hợp rồi chèn vào mục Use Case, ERD và Deployment Diagram.')
for heading, body in SECTIONS:
    report.add_page_break()
    report.add_heading(heading, level=1)
    for paragraph in body.split('\n'):
        if paragraph.strip(): report.add_paragraph(paragraph.strip())
    if heading.startswith('7.'):
        report.add_paragraph('Browser → Application (Express/EJS + REST API) → PostgreSQL và MinIO. Database và MinIO dùng volumes riêng.')
    if heading.startswith('8.'):
        table = report.add_table(rows=1, cols=2); table.style = 'Light Shading Accent 2'
        table.rows[0].cells[0].text = 'Nhóm bảng'; table.rows[0].cells[1].text = 'Mục đích'
        for a,b in [('Danh mục','categories, brands, suppliers'),('Bán hàng','products, product_batches, orders, order_items'),('Người dùng','users, audit_logs'),('Lưu trữ','file_metadata + object trong MinIO'),('Kho','stock_movements')]:
            cells = table.add_row().cells; cells[0].text = a; cells[1].text = b
    if heading.startswith('16.'):
        for item in ['application — cổng 3000','database — PostgreSQL 5432 nội bộ, volume database-data','minio — API 9000, Console 9001, volume minio-data']:
            report.add_paragraph(item, style='List Bullet')
report_path = DOCS / 'BaoCao_DoAn_Mau.docx'
report.save(report_path)

# PDF report: one section per page to make a 20+ page editable-ready draft.
font_regular = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
font_bold = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
pdfmetrics.registerFont(TTFont('DV', font_regular)); pdfmetrics.registerFont(TTFont('DV-Bold', font_bold))
pdf_path = DOCS / 'BaoCao_DoAn_Mau.pdf'
styles_pdf = getSampleStyleSheet()
styles_pdf.add(ParagraphStyle(name='CoverTitle', fontName='DV-Bold', fontSize=25, leading=32, alignment=TA_CENTER, textColor=colors.HexColor('#A65F61'), spaceAfter=18))
styles_pdf.add(ParagraphStyle(name='CoverSub', fontName='DV-Bold', fontSize=14, leading=20, alignment=TA_CENTER, textColor=colors.HexColor('#4B403A'), spaceAfter=12))
styles_pdf.add(ParagraphStyle(name='Chapter', fontName='DV-Bold', fontSize=16, leading=22, textColor=colors.HexColor('#A65F61'), spaceAfter=15))
styles_pdf.add(ParagraphStyle(name='BodyVN', fontName='DV', fontSize=9.5, leading=15, alignment=TA_LEFT, textColor=colors.HexColor('#39322E'), spaceAfter=9))
styles_pdf.add(ParagraphStyle(name='SmallVN', fontName='DV', fontSize=8, leading=12, textColor=colors.HexColor('#6D625C'), spaceAfter=6))
def footer(canvas, doc):
    canvas.saveState(); canvas.setStrokeColor(colors.HexColor('#EDE4DE')); canvas.line(18*mm, 14*mm, 192*mm, 14*mm)
    canvas.setFont('DV', 7); canvas.setFillColor(colors.HexColor('#948981'))
    canvas.drawString(18*mm, 9*mm, 'Lumière Cosmetic · Báo cáo mẫu đồ án Điện toán đám mây')
    canvas.drawRightString(192*mm, 9*mm, f'Trang {doc.page}')
    canvas.restoreState()
story = [Spacer(1, 35*mm), Paragraph('LUMIÈRE COSMETIC', styles_pdf['CoverTitle']), Paragraph('BÁO CÁO ĐỒ ÁN CUỐI KỲ', styles_pdf['CoverSub']), Paragraph('HỌC PHẦN: ĐIỆN TOÁN ĐÁM MÂY', styles_pdf['CoverSub']), Spacer(1, 20*mm)]
for line in ['Đề tài: Website quản lý mỹ phẩm mini đóng gói bằng Docker Compose','Sinh viên: [Bổ sung họ tên thành viên]','Mã sinh viên: [Bổ sung]','Lớp: [Bổ sung]','Giảng viên hướng dẫn: [Bổ sung]','Đơn vị: [Bổ sung trường / khoa]','Năm học: 2026']:
    story.append(Paragraph(line, ParagraphStyle(name='CoverLine'+str(len(story)), parent=styles_pdf['BodyVN'], alignment=TA_CENTER, fontSize=10, leading=18)))
story += [Spacer(1, 10*mm), Paragraph('TÓM TẮT', styles_pdf['CoverSub']), Paragraph('Ứng dụng web quản lý cửa hàng mỹ phẩm mini với Express, PostgreSQL, MinIO và Docker Compose; có CRUD, quản lý lô/hạn dùng, bán hàng, upload/download, phân quyền và kiểm chứng dữ liệu bền vững.', styles_pdf['BodyVN']), Spacer(1, 10*mm), Paragraph('BẢN MẪU — CẦN BỔ SUNG TÊN THÀNH VIÊN, ẢNH CHỤP VÀ KẾT QUẢ CHẠY THỰC TẾ TRƯỚC KHI NỘP', styles_pdf['SmallVN']), PageBreak()]
for heading, body in SECTIONS:
    if heading == 'Tóm tắt':
        continue
    story.append(Paragraph(heading, styles_pdf['Chapter']))
    story.append(Paragraph(body.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'), styles_pdf['BodyVN']))
    extras = {
      '7. Kiến trúc tổng thể': 'Luồng: Browser → Application container (Node.js/Express/EJS) → PostgreSQL và MinIO. Docker Compose tạo network và hai named volumes.',
      '8. Thiết kế database': 'Bảng nghiệp vụ được liên kết bằng khóa ngoại. Các ràng buộc CHECK và UNIQUE bảo vệ tính hợp lệ; chỉ metadata tệp nằm trong SQL, còn nội dung file nằm trong MinIO.',
      '16. Dockerfile và Compose': 'Service application publish 3000; database dùng 5432 nội bộ; MinIO publish 9000 và 9001. Volume database-data và minio-data cần tồn tại sau lệnh docker compose down.',
      '17. Liên hệ VPC/subnet': 'Docker bridge network là một cơ chế mạng container có phạm vi giới hạn; không thay thế đầy đủ tính năng VPC/subnet, route table và security group của cloud.',
      '18. Kiểm thử': 'Bằng chứng cần chèn: ảnh trang dashboard, CRUD, file trong MinIO Console, Swagger/Postman và trạng thái dữ liệu trước/sau khi down/up.'
    }
    if heading in extras: story.append(Paragraph(extras[heading], styles_pdf['BodyVN']))
    story.append(Spacer(1, 6*mm)); story.append(PageBreak())
SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=17*mm, bottomMargin=19*mm, title='Báo cáo Lumière Cosmetic').build(story, onFirstPage=footer, onLaterPages=footer)

# PowerPoint defense deck
prs = Presentation(); prs.slide_width = PptInches(13.333); prs.slide_height = PptInches(7.5)
BG = PptRGB(250,248,246); ROSE = PptRGB(166,95,97); DARK = PptRGB(48,42,38); MUTED = PptRGB(126,116,108); PALE = PptRGB(248,239,236); GREEN = PptRGB(231,241,232)
def set_bg(slide, color=BG):
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = color
def add_text(slide, text, x, y, w, h, size=20, color=DARK, bold=False, font='Aptos'):
    box = slide.shapes.add_textbox(PptInches(x),PptInches(y),PptInches(w),PptInches(h)); tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.name = font; p.font.size = PptPt(size); p.font.bold = bold; p.font.color.rgb = color
    return box
def add_title(slide, title, subtitle=''):
    add_text(slide, title, .7, .45, 11.8, .55, 27, ROSE, True, 'Aptos Display')
    if subtitle: add_text(slide, subtitle, .73, 1.08, 11.5, .38, 11, MUTED)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,PptInches(.72),PptInches(1.53),PptInches(.75),PptInches(.045)); shape.fill.solid(); shape.fill.fore_color.rgb=ROSE; shape.line.fill.background()
def add_bullets(slide, bullets, x=.9, y=1.9, w=11.5, h=4.8, size=19):
    box=slide.shapes.add_textbox(PptInches(x),PptInches(y),PptInches(w),PptInches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    for idx, item in enumerate(bullets):
        p=tf.paragraphs[0] if idx==0 else tf.add_paragraph(); p.text=item; p.level=0; p.font.name='Aptos'; p.font.size=PptPt(size); p.font.color.rgb=DARK; p.space_after=PptPt(16); p.bullet=True
    return box
slide=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide, PptRGB(248,239,236))
add_text(slide,'LUMIÈRE',.85,1.0,5,.6,28,ROSE,True,'Aptos Display'); add_text(slide,'COSMETIC',.87,1.58,4,.4,13,MUTED,True)
add_text(slide,'WEBSITE QUẢN LÝ\nMỸ PHẨM MINI',.85,2.45,7.2,1.65,34,DARK,True,'Aptos Display')
add_text(slide,'Đồ án cuối kỳ · Học phần Điện toán đám mây',.9,4.35,7,.4,16,MUTED)
add_text(slide,'Sinh viên: [Tên nhóm]  |  Giảng viên: [Tên giảng viên]',.9,6.55,9,.35,11,MUTED)
# slides content
slides = [
('Bài toán và lý do chọn đề tài',['Dữ liệu sản phẩm, khách hàng và đơn hàng thường nằm ở nhiều nơi.','Khó theo dõi tồn kho theo lô và hạn sử dụng.','Ảnh sản phẩm/chứng từ cần lưu trữ tách khỏi database.','Mục tiêu: một ứng dụng web chạy được bằng Docker Compose.']),
('Mục tiêu và phạm vi',['CRUD sản phẩm, danh mục, nhãn hàng, nhà cung cấp, khách hàng.','Nhập lô, theo dõi hạn sử dụng, bán hàng và in hóa đơn.','Upload/download file qua MinIO.','Báo cáo doanh thu, tồn kho thấp và sản phẩm bán chạy.']),
('Kiến trúc hệ thống',['Browser → Node.js / Express / EJS.','Backend → PostgreSQL để lưu nghiệp vụ và metadata.','Backend → MinIO để lưu nội dung file.','Docker Compose kết nối ba service bằng network riêng.']),
('Công nghệ sử dụng',['Node.js 22 + Express REST API.','EJS, CSS, JavaScript và Bootstrap Icons.','PostgreSQL 16, foreign keys, constraints và indexes.','MinIO Object Storage, Dockerfile và Docker Compose.']),
('Mô hình dữ liệu',['13 bảng nghiệp vụ, người dùng, kho, metadata và audit.','Product liên kết Category, Brand, Supplier.','Order có nhiều Order Item; lô hàng lưu số lượng còn lại.','File metadata trỏ đến object key trong bucket MinIO.']),
('Chức năng giao diện',['Dashboard với KPI và cảnh báo.','Danh sách có tìm kiếm, CRUD và trang chi tiết.','Form nhập lô, tạo đơn và in hóa đơn.','Kho tệp có xem/tải xuống và thông báo kết quả.']),
('REST API và phân quyền',['REST API cho sản phẩm, danh mục, nhãn hàng, nhà cung cấp, khách hàng.','Swagger UI / OpenAPI và Postman Collection.','Administrator quản trị dữ liệu và tài khoản.','Staff thao tác nghiệp vụ theo quyền; mật khẩu bcrypt hash.']),
('Upload file qua MinIO',['Browser gửi multipart/form-data đến backend.','Backend kiểm tra dung lượng, phần mở rộng và MIME type.','MinIO lưu nội dung file; PostgreSQL chỉ lưu metadata.','File được xem/tải qua route yêu cầu đăng nhập.']),
('Docker Compose và persistence',['Service: application, database, minio.','Named volumes: database-data và minio-data.','Network: cosmetic-network, dùng DNS service.','Kiểm chứng bằng down rồi up -d, không dùng down -v.']),
('Kiểm thử và demo',['CRUD ít nhất 3 thực thể và tìm kiếm.','Tạo đơn, kiểm tra tồn kho và xử lý bán vượt tồn.','Upload/download file; xem object trong MinIO Console.','Gọi API qua Swagger/Postman; kiểm tra dữ liệu sau restart.']),
('Liên hệ lý thuyết điện toán đám mây',['Docker network cô lập giao tiếp container, có điểm tương đồng mục tiêu với VPC/subnet.','Named volume tách dữ liệu khỏi vòng đời container.','Object Storage phù hợp lưu nội dung file không cấu trúc.','Compose giúp tái tạo môi trường demo nhất quán.']),
('Hạn chế và hướng phát triển',['Bổ sung CSRF, rate limiting, session store bền vững và backup.','Hoàn thiện hủy/hoàn đơn, khuyến mãi và báo cáo lợi nhuận.','Tạo MinIO policy least-privilege cho triển khai thật.','Bổ sung ảnh chụp, kết quả test và phân công nhóm trước khi nộp.'])]
for title, bullets in slides:
    slide=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide); add_title(slide,title); add_bullets(slide,bullets,size=18)
    add_text(slide,'LUMIÈRE COSMETIC  ·  ĐỒ ÁN ĐIỆN TOÁN ĐÁM MÂY',.75,7.12,8,.2,8,MUTED,True)
slide=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide,PptRGB(248,239,236)); add_text(slide,'CẢM ƠN',.8,2.35,11.8,1.0,38,ROSE,True,'Aptos Display'); add_text(slide,'Câu hỏi & thảo luận',.85,3.55,10,.5,22,DARK); add_text(slide,'[Tên thành viên] · [Lớp]',.88,6.45,8,.3,11,MUTED)
ppt_path = DOCS / 'Slide_BaoVe_Mau.pptx'; prs.save(ppt_path)
print(f'Created: {report_path}')
print(f'Created: {pdf_path}')
print(f'Created: {ppt_path}')
